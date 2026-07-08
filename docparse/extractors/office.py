"""DOCX and XLSX extractors (heading levels, list items, three-view tables)."""
from __future__ import annotations

import os
import re

from docparse.router import register
from docparse.models import Element, ExtractionConfig, ExtractionResult
from docparse.serialize import serialize_table

_HEADING_RE = re.compile(r"heading\s*(\d+)", re.IGNORECASE)


class DocxExtractor:
    doc_types = ("docx",)

    def extract(self, path: str, cfg: ExtractionConfig) -> ExtractionResult:
        import docx
        from docx.oxml.ns import qn
        from docx.table import Table
        from docx.text.paragraph import Paragraph

        result = ExtractionResult(source=path, doc_type="docx",
                                  mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        document = docx.Document(path)
        order = 0
        pending_caption: Element | None = None
        img_n = 0
        if cfg.assets_dir and cfg.extract_figures:
            os.makedirs(cfg.assets_dir, exist_ok=True)
        for child in document.element.body.iterchildren():
            if child.tag == qn("w:p"):
                para = Paragraph(child, document)
                # embedded images: any drawing/blip in the paragraph -> figure element
                if cfg.extract_figures:
                    for blip in child.findall(".//" + qn("a:blip")):
                        rid = blip.get(qn("r:embed"))
                        rel = document.part.rels.get(rid) if rid else None
                        if rel is None or "image" not in rel.reltype:
                            continue
                        img_n += 1
                        visual = {"rel_id": rid, "content_type": rel.target_part.content_type}
                        if cfg.assets_dir:
                            ext = (rel.target_part.content_type or "image/png").split("/")[-1]
                            fp = os.path.join(cfg.assets_dir, f"docx_fig{img_n}.{ext}")
                            with open(fp, "wb") as fh:
                                fh.write(rel.target_part.blob)
                            visual["image_path"] = fp
                        order += 1
                        fig = Element(kind="figure", text="", order=order,
                                      extractor="python-docx", visual_ref=visual)
                        result.elements.append(fig)
                        if pending_caption is not None:       # "Figure N:" directly above
                            pending_caption.parent_id = fig.id
                            pending_caption = None
                text = para.text.strip()
                if not text:
                    continue
                style = (para.style.name or "") if para.style else ""
                sl = style.lower()
                kind, level = "narrative", None
                if sl.startswith("title"):
                    kind, level = "title", 0
                elif (m := _HEADING_RE.match(sl)):
                    kind, level = "heading", min(int(m.group(1)), 6)
                elif sl.startswith("heading"):
                    kind, level = "heading", 1
                elif "list" in sl or child.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None \
                        or (child.find(qn("w:pPr")) is not None
                            and child.find(qn("w:pPr")).find(qn("w:numPr")) is not None):
                    kind = "list_item"
                elif sl.startswith("caption"):
                    kind = "caption"
                order += 1
                el = Element(kind=kind, text=text, order=order, level=level,
                             extractor="python-docx", metadata={"style": style or None})
                if kind == "caption" and result.elements \
                        and result.elements[-1].kind == "figure" \
                        and result.elements[-1].parent_id is None:
                    el.parent_id = result.elements[-1].id     # caption right below its image
                result.elements.append(el)
                pending_caption = el if kind == "caption" else None
            elif child.tag == qn("w:tbl"):
                table = Table(child, document)
                rows = [[c.text for c in row.cells] for row in table.rows]
                text, extras, structured, trunc = serialize_table(
                    rows, cfg.table_serializations, cfg.table_max_rows)
                if not text:
                    continue
                order += 1
                meta = {"rows": len(rows), "truncated": trunc}
                if extras:
                    meta["serializations"] = extras
                el = Element(kind="table", text=text, order=order, extractor="python-docx",
                             structured=structured, metadata=meta)
                result.elements.append(el)
                if pending_caption is not None:      # "Table 1: ..." directly above
                    pending_caption.parent_id = el.id
                    pending_caption = None
        return result


class XlsxExtractor:
    doc_types = ("xlsx",)

    def extract(self, path: str, cfg: ExtractionConfig) -> ExtractionResult:
        from openpyxl import load_workbook

        result = ExtractionResult(source=path, doc_type="xlsx",
                                  mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        wb = load_workbook(path, read_only=True, data_only=True)
        sheet_names = list(wb.sheetnames)
        order = 0
        for name in sheet_names:
            ws = wb[name]
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            rows = [r for r in rows if any(c is not None and str(c).strip() for c in r)]
            if not rows:
                continue
            # a formula workbook saved without cached values yields Nones under data_only
            if all(all(c is None for c in r) for r in rows):
                result.warnings.append(f"sheet '{name}': no cached formula values (data_only)")
                continue
            text, extras, structured, trunc = serialize_table(
                rows, cfg.table_serializations, cfg.table_max_rows)
            order += 1
            meta = {"sheet": name, "rows": len(rows),
                    "cols": max(len(r) for r in rows), "truncated": trunc}
            if extras:
                meta["serializations"] = extras
            result.elements.append(Element(kind="table", text=text, order=order,
                                           extractor="openpyxl", structured=structured,
                                           metadata=meta))
        result.stats["sheets"] = sheet_names
        wb.close()
        return result


register(DocxExtractor())
register(XlsxExtractor())
