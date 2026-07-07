"""Build a small synthetic golden corpus (PDF + DOCX) with ground truth for bench."""
from __future__ import annotations

import json
from pathlib import Path

LOREM = ("This paragraph carries enough ordinary running text to be treated as narrative "
         "content by any sane extractor and to make language detection possible.")


def make_pdf(path: Path) -> dict:
    import fitz
    doc = fitz.open()
    W, H = 595, 842
    gt = {"headings": [], "reading_order": [], "tables": []}

    for pno in range(3):
        page = doc.new_page(width=W, height=H)
        page.insert_text((200, 30), "ACME Corp — Annual Report", fontsize=8)      # running header
        page.insert_text((280, H - 25), f"Page {pno + 1}", fontsize=8)            # running footer

        if pno == 0:
            page.insert_text((72, 90), "Quarterly Results", fontsize=20)
            gt["headings"].append("Quarterly Results")
            gt["reading_order"].append("Quarterly Results")
            # two columns: left column first, then right
            left = f"Left column opening statement for page one. {LOREM}"
            right = f"Right column continuation text for page one. {LOREM}"
            page.insert_textbox((72, 120, 285, 420), left, fontsize=10)
            # right column starts *higher* than the left: naive y-sort reads it first,
            # column-aware XY-cut must still read left column first
            page.insert_textbox((310, 112, 523, 420), right, fontsize=10)
            gt["reading_order"] += ["Left column opening statement for page one",
                                    "Right column continuation text for page one"]
            page.insert_text((72, 460), "Methodology", fontsize=15)
            gt["headings"].append("Methodology")
            gt["reading_order"].append("Methodology")
            page.insert_textbox((72, 480, 523, 560),
                                f"Methodology details follow the headline. {LOREM}", fontsize=10)
            gt["reading_order"].append("Methodology details follow the headline")
        elif pno == 1:
            page.insert_text((72, 90), "Financial Table", fontsize=15)
            gt["headings"].append("Financial Table")
            rows = [["Region", "Revenue", "Growth"],
                    ["North", "120", "4%"],
                    ["South", "95", "7%"]]
            x0, y0, cw, rh = 72, 120, 120, 24
            for r, row in enumerate(rows):
                for c, cell in enumerate(row):
                    page.insert_text((x0 + c * cw + 6, y0 + r * rh + 16), cell, fontsize=10)
            for r in range(len(rows) + 1):
                page.draw_line((x0, y0 + r * rh), (x0 + 3 * cw, y0 + r * rh))
            for c in range(4):
                page.draw_line((x0 + c * cw, y0), (x0 + c * cw, y0 + len(rows) * rh))
            gt["tables"].append({"page": 2, "rows": rows})
            page.insert_text((72, 230), "Table 1: Revenue by region.", fontsize=9)
            page.insert_textbox((72, 260, 523, 340),
                                f"Discussion of the revenue table. {LOREM}", fontsize=10)
            gt["reading_order"] += ["Financial Table", "Discussion of the revenue table"]
        else:
            page.insert_text((72, 90), "Outlook", fontsize=15)
            gt["headings"].append("Outlook")
            page.insert_textbox((72, 120, 523, 260),
                                f"Closing outlook narrative on the final page. {LOREM}", fontsize=10)
            gt["reading_order"] += ["Outlook", "Closing outlook narrative on the final page"]

    doc.save(str(path))
    doc.close()
    return gt


def make_docx(path: Path) -> dict:
    import docx
    d = docx.Document()
    d.add_heading("Project Plan", level=1)
    d.add_heading("Scope", level=2)
    d.add_paragraph(f"The scope paragraph explains boundaries. {LOREM}")
    t = d.add_table(rows=2, cols=2)
    t.rows[0].cells[0].text, t.rows[0].cells[1].text = "Task", "Owner"
    t.rows[1].cells[0].text, t.rows[1].cells[1].text = "Parsing", "Alice"
    d.save(str(path))
    return {"headings": ["Project Plan", "Scope"],
            "reading_order": ["Project Plan", "Scope", "The scope paragraph explains boundaries"],
            "tables": [{"rows": [["Task", "Owner"], ["Parsing", "Alice"]]}]}


def build(out_dir: str | Path) -> Path:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    pdf = out / "report.pdf"
    (out / "report.pdf.gt.json").write_text(json.dumps(make_pdf(pdf), indent=2), encoding="utf-8")
    dx = out / "plan.docx"
    (out / "plan.docx.gt.json").write_text(json.dumps(make_docx(dx), indent=2), encoding="utf-8")
    return out


if __name__ == "__main__":
    import sys
    print(build(sys.argv[1] if len(sys.argv) > 1 else "golden"))
